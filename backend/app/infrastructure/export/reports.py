"""Xuất báo cáo (UC17) ra Excel (.xlsx) và PDF. Font DejaVu để in đúng tiếng Việt có dấu."""

import io
import pathlib

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.application.use_cases.reports import Report

FONTS = pathlib.Path(__file__).resolve().parents[2] / "assets" / "fonts"
MONEY = '#,##0" đ"'


def _title(ws, text: str, rep: Report) -> None:
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Từ ngày {rep.date_from:%d/%m/%Y} đến ngày {rep.date_to:%d/%m/%Y}"


def _header(ws, row: int, cols: list[str]) -> None:
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center")


def to_xlsx(rep: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Doanh thu"
    _title(ws, "BÁO CÁO DOANH THU THEO TUYẾN", rep)
    _header(ws, 4, ["Tuyến", "Doanh thu", "Đã hoàn", "Thực thu", "Vé bán", "Vé hủy"])
    r = 5
    for x in rep.by_route.values():
        ws.append(
            [
                x.route,
                float(x.gross),
                float(x.refunded),
                float(x.net),
                x.tickets_sold,
                x.tickets_cancelled,
            ]
        )
        r += 1
    ws.append(
        [
            "TỔNG",
            float(rep.gross),
            float(rep.refunded),
            float(rep.net),
            rep.tickets_sold,
            rep.tickets_cancelled,
        ]
    )
    for c in ws[r]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=5, max_row=r, min_col=2, max_col=4):
        for c in row:
            c.number_format = MONEY
    ws.append([])
    ws.append(["Doanh thu thực thu theo kênh"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for ch, v in rep.net_by_channel.items():
        ws.append([{"TRUC_TUYEN": "Trực tuyến", "TAI_QUAY": "Tại quầy"}.get(ch, ch), float(v)])
        ws.cell(row=ws.max_row, column=2).number_format = MONEY
    ws.append([])
    ws.append(["Tỷ lệ hủy vé", rep.cancel_rate])
    ws.cell(row=ws.max_row, column=2).number_format = "0.0%"

    fill = wb.create_sheet("Lấp đầy")
    _title(fill, "TỶ LỆ LẤP ĐẦY CHỖ NGỒI", rep)
    _header(fill, 4, ["Tuyến", "Ghế đã bán", "Tổng ghế", "Tỷ lệ lấp đầy"])
    for f in rep.fill_by_route.values():
        fill.append([f.route, f.sold, f.total, f.rate])
        fill.cell(row=fill.max_row, column=4).number_format = "0.0%"
    fill.append(
        [
            "Toàn kỳ",
            sum(f.sold for f in rep.fill_by_route.values()),
            sum(f.total for f in rep.fill_by_route.values()),
            rep.occupancy,
        ]
    )
    fill.cell(row=fill.max_row, column=4).number_format = "0.0%"

    for sheet in (ws, fill):
        for col in range(1, 7):
            sheet.column_dimensions[get_column_letter(col)].width = 28 if col == 1 else 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _money(v) -> str:
    return f"{float(v):,.0f} đ".replace(",", ".")


def to_pdf(rep: Report) -> bytes:
    pdf = FPDF(orientation="L", format="A4")
    pdf.add_font("DejaVu", "", str(FONTS / "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", str(FONTS / "DejaVuSans-Bold.ttf"))
    pdf.add_page()
    pdf.set_font("DejaVu", "B", 16)
    pdf.cell(0, 10, "BÁO CÁO DOANH THU & LẤP ĐẦY", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 11)
    pdf.cell(
        0,
        7,
        f"Từ ngày {rep.date_from:%d/%m/%Y} đến ngày {rep.date_to:%d/%m/%Y}",
        new_x="LMARGIN",
        new_y="NEXT",
    )
    pdf.cell(
        0,
        7,
        f"Doanh thu {_money(rep.gross)}   Đã hoàn {_money(rep.refunded)}   Thực thu {_money(rep.net)}",
        new_x="LMARGIN",
        new_y="NEXT",
    )
    pdf.cell(
        0,
        7,
        f"Vé bán {rep.tickets_sold}   Vé hủy {rep.tickets_cancelled}   Tỷ lệ hủy {rep.cancel_rate:.1%}   Lấp đầy {rep.occupancy:.1%}",
        new_x="LMARGIN",
        new_y="NEXT",
    )
    pdf.ln(4)

    def table(headers: list[str], widths: list[int], rows: list[list[str]]) -> None:
        pdf.set_font("DejaVu", "B", 10)
        pdf.set_fill_color(37, 99, 235)
        pdf.set_text_color(255, 255, 255)
        for h, w in zip(headers, widths, strict=True):
            pdf.cell(w, 8, h, border=1, fill=True, align="C")
        pdf.ln()
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("DejaVu", "", 10)
        for row in rows:
            for i, (v, w) in enumerate(zip(row, widths, strict=True)):
                pdf.cell(w, 7, v, border=1, align="L" if i == 0 else "R")
            pdf.ln()
        pdf.ln(4)

    table(
        ["Tuyến", "Doanh thu", "Đã hoàn", "Thực thu", "Vé bán", "Vé hủy"],
        [90, 40, 40, 40, 25, 25],
        [
            [
                x.route,
                _money(x.gross),
                _money(x.refunded),
                _money(x.net),
                str(x.tickets_sold),
                str(x.tickets_cancelled),
            ]
            for x in rep.by_route.values()
        ]
        or [["(không có dữ liệu trong kỳ)", "", "", "", "", ""]],
    )
    table(
        ["Tuyến", "Ghế đã bán", "Tổng ghế", "Tỷ lệ lấp đầy"],
        [90, 40, 40, 40],
        [[f.route, str(f.sold), str(f.total), f"{f.rate:.1%}"] for f in rep.fill_by_route.values()]
        or [["(không có chuyến trong kỳ)", "", "", ""]],
    )
    return bytes(pdf.output())
